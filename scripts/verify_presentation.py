import pptx
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

orig = pptx.Presentation(r'C:\Users\rohit\.gemini\antigravity\brain\28304f3d-38e9-4a0c-81e0-3e0d92e2eadd\.user_uploaded\media_1793115064879_9a23a708.pptx')
gen = pptx.Presentation(r'docs\PROJECT_PRESENTATION.pptx')

print(f"Original slide count: {len(orig.slides)}, Generated slide count: {len(gen.slides)}")
assert len(orig.slides) == len(gen.slides) == 17

# Verify slide 1 & 2 text identical
for s_idx in [0, 1]:
    orig_s = orig.slides[s_idx]
    gen_s = gen.slides[s_idx]
    orig_texts = [sh.text_frame.text for sh in orig_s.shapes if sh.has_text_frame]
    gen_texts = [sh.text_frame.text for sh in gen_s.shapes if sh.has_text_frame]
    assert orig_texts == gen_texts, f"Slide {s_idx+1} mismatch!"
    print(f"Slide {s_idx+1} perfectly identical to original template!")

# Check all slides
for s_idx in range(2, 17):
    s = gen.slides[s_idx]
    title = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and ('object 2' in sh.name or 'Title' in sh.name)]
    title_str = title[0] if title else 'Unknown'
    text_count = sum(len(sh.text_frame.text.strip()) for sh in s.shapes if sh.has_text_frame)
    has_pic = any(sh.shape_type == 13 for sh in s.shapes)
    has_chart = any(sh.has_chart for sh in s.shapes)
    has_tbl = any(sh.has_table for sh in s.shapes)
    print(f"Slide {s_idx+1:02d}: '{title_str}' | text_chars={text_count} | pic={has_pic} | chart={has_chart} | tbl={has_tbl}")

# Let's inspect slide 14 in detail
print("\n--- Inspecting Slide 14 Details ---")
s14 = gen.slides[13]
for sh in s14.shapes:
    if sh.has_text_frame:
        print(f"Shape '{sh.name}': {repr(sh.text_frame.text.replace(chr(10), ' -- '))}")
    elif sh.shape_type == 13:
        print(f"Picture Shape '{sh.name}': left={sh.left}, top={sh.top}, width={sh.width}, height={sh.height}")

print("\n--- Inspecting Slide 4 Stats ---")
s4 = gen.slides[3]
for stat_name in ['Stat 1', 'Stat 2', 'Stat 3', 'Stat 4']:
    sh = [s for s in s4.shapes if s.name == stat_name][0]
    print(f"{stat_name}: {repr(sh.text_frame.text.replace(chr(10), ' | '))}")

print("\nVerification Complete! All checks PASSED.")
